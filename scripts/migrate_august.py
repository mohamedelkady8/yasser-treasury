#!/usr/bin/env python3
"""ترحيل بيانات ملف «AUGUST - 26.xlsx» إلى قاعدة بيانات Supabase.

يقرأ الملف بـ openpyxl، يرحّل القوائم المرجعية أولًا، ثم القيود بربطها
بالمفاتيح الأجنبية بالاسم، ثم يطابق الإجماليات مع الإكسل.

الاستخدام:
    python3 scripts/migrate_august.py            # ترحيل
    python3 scripts/migrate_august.py --verify   # مطابقة الإجماليات فقط
"""

from __future__ import annotations

import json
import os
import sys
import urllib.error
import urllib.request
import warnings
from datetime import date, datetime
from pathlib import Path

import openpyxl

warnings.filterwarnings("ignore")

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _config import ANON_KEY, EMAIL, PASSWORD, ROOT, SUPABASE_URL, require  # noqa: E402

XLSX = Path(os.environ.get("AUGUST_XLSX") or ROOT.parent / "AUGUST - 26.xlsx")

# القيم الافتراضية للقيود ناقصة البيانات في الملف الأصلي
DEFAULT_REVENUE_DATE = date(2026, 8, 1)
UNKNOWN_REVENUE_TYPE = "إيراد غير معلوم"
OPENING_REVENUE_TYPE = "رصيد أفتتاحي"

MOVEMENT_MAP = {
    "تشغيلي": "operational",
    "تحويل داخلي": "internal_transfer",
    "صرف عهدة": "custody_out",
    "مصروف من العهدة": "custody_expense",
    "مرتجع عهدة": "custody_return",
}
RECON_MAP = {"تمت": "reconciled", "لم تتم": "pending"}


# --------------------------------------------------------------------------- HTTP


class Api:
    def __init__(self, url: str, key: str):
        self.url = url
        self.key = key
        self.token = key

    def _req(self, method: str, path: str, body=None, headers=None):
        url = path if path.startswith("http") else f"{self.url}{path}"
        data = json.dumps(body).encode() if body is not None else None
        req = urllib.request.Request(url, data=data, method=method)
        req.add_header("apikey", self.key)
        req.add_header("Authorization", f"Bearer {self.token}")
        req.add_header("Content-Type", "application/json")
        for k, v in (headers or {}).items():
            req.add_header(k, v)
        try:
            with urllib.request.urlopen(req) as resp:
                raw = resp.read().decode()
                return json.loads(raw) if raw else None
        except urllib.error.HTTPError as exc:
            detail = exc.read().decode()
            raise SystemExit(f"{method} {url} -> {exc.code}\n{detail}") from exc

    def login(self, email: str, password: str):
        res = self._req(
            "POST",
            "/auth/v1/token?grant_type=password",
            {"email": email, "password": password},
        )
        self.token = res["access_token"]
        return res

    def select(self, table: str, query: str = "select=*"):
        return self._req("GET", f"/rest/v1/{table}?{query}") or []

    def insert(self, table: str, rows: list, upsert=False):
        if not rows:
            return []
        prefer = "return=representation"
        if upsert:
            prefer += ",resolution=merge-duplicates"
        out = []
        for i in range(0, len(rows), 200):
            chunk = rows[i : i + 200]
            res = self._req(
                "POST", f"/rest/v1/{table}", chunk, {"Prefer": prefer}
            )
            out.extend(res or [])
        return out

    def rpc(self, fn: str, args: dict | None = None):
        return self._req("POST", f"/rest/v1/rpc/{fn}", args or {})

    def delete_all(self, table: str, pk: str = "id"):
        self._req("DELETE", f"/rest/v1/{table}?{pk}=not.is.null")


# --------------------------------------------------------------------------- Excel


def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def as_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def read_workbook():
    if not XLSX.exists():
        raise SystemExit(f"لم أجد الملف: {XLSX}")
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    lists_ws, entries_ws = wb["القوائم"], wb["القيود"]

    def column(ws, col: str, first_row=4, last_row=200):
        out = []
        for r in range(first_row, last_row):
            v = clean(ws[f"{col}{r}"].value)
            if v is not None and v not in out:
                out.append(v)
        return out

    def paired(ws, name_col: str, value_col: str):
        out = {}
        for r in range(4, 200):
            name = clean(ws[f"{name_col}{r}"].value)
            if name is None:
                continue
            val = ws[f"{value_col}{r}"].value
            out[name] = float(val) if isinstance(val, (int, float)) else 0.0
        return out

    lists = {
        "accounts": column(lists_ws, "A"),
        "ledgers": column(lists_ws, "C"),
        "cost_centers": column(lists_ws, "E"),
        "revenue_types": column(lists_ws, "H"),
        "banks": column(lists_ws, "J"),
    }
    lists["custody_opening"] = paired(lists_ws, "E", "F")
    lists["bank_opening"] = paired(lists_ws, "J", "K")

    revenue, expense = [], []
    for r in range(8, entries_ws.max_row + 1):
        amount = entries_ws.cell(r, 2).value
        if isinstance(amount, (int, float)) and amount:
            revenue.append(
                {
                    "row": r,
                    "amount": round(float(amount), 2),
                    "description": clean(entries_ws.cell(r, 3).value),
                    "entry_date": as_date(entries_ws.cell(r, 4).value),
                    "revenue_type": clean(entries_ws.cell(r, 5).value),
                    "bank": clean(entries_ws.cell(r, 6).value),
                    "note": clean(entries_ws.cell(r, 7).value),
                    "movement": clean(entries_ws.cell(r, 8).value),
                    "recon": clean(entries_ws.cell(r, 9).value),
                }
            )
        amount = entries_ws.cell(r, 13).value
        if isinstance(amount, (int, float)) and amount:
            expense.append(
                {
                    "row": r,
                    "entry_code": clean(entries_ws.cell(r, 12).value),
                    "amount": round(float(amount), 2),
                    "description": clean(entries_ws.cell(r, 14).value),
                    "entry_date": as_date(entries_ws.cell(r, 15).value),
                    "account": clean(entries_ws.cell(r, 16).value),
                    "ledger": clean(entries_ws.cell(r, 17).value),
                    "cost_center": clean(entries_ws.cell(r, 18).value),
                    "bank": clean(entries_ws.cell(r, 19).value),
                    "note": clean(entries_ws.cell(r, 20).value),
                    "document": clean(entries_ws.cell(r, 21).value),
                    "movement": clean(entries_ws.cell(r, 22).value),
                    "recon": clean(entries_ws.cell(r, 23).value),
                }
            )
    return lists, revenue, expense


# --------------------------------------------------------------------------- migrate


def migrate(api: Api):
    lists, revenue, expense = read_workbook()
    print(
        f"الملف: {len(revenue)} قيد إيراد، {len(expense)} قيد مصروف، "
        f"{sum(len(v) for k, v in lists.items() if isinstance(v, list))} بند قوائم"
    )

    # 1) القوائم المرجعية — إدخال ما هو ناقص فقط ليكون السكربت قابلًا للتكرار
    def seed(table: str, rows: list[dict]):
        seen, unique = set(), []
        for row in rows:
            key = row["name"].strip()
            if key in seen:
                continue
            seen.add(key)
            row["name"] = key
            unique.append(row)
        have = {r["name"].strip() for r in api.select(table, "select=name")}
        missing = [r for r in unique if r["name"] not in have]
        api.insert(table, missing)
        return len(unique), len(missing)

    seed("accounts", [{"name": n} for n in lists["accounts"]])
    seed("ledgers", [{"name": n} for n in lists["ledgers"]])
    seed(
        "cost_centers",
        [
            {"name": n, "custody_opening": lists["custody_opening"].get(n, 0)}
            for n in lists["cost_centers"]
        ],
    )
    seed(
        "revenue_types",
        [
            {"name": n, "is_opening": n.strip() == OPENING_REVENUE_TYPE}
            for n in lists["revenue_types"] + [UNKNOWN_REVENUE_TYPE]
        ],
    )
    seed(
        "banks",
        [
            {
                "name": n,
                "opening_balance": lists["bank_opening"].get(n, 0),
                "is_usd": "دولار" in n or "DOLLAR" in n.upper(),
                "sort_order": i,
            }
            for i, n in enumerate(lists["banks"])
        ],
    )

    ids = {
        table: {row["name"].strip(): row["id"] for row in api.select(table, "select=id,name")}
        for table in ("accounts", "ledgers", "cost_centers", "revenue_types", "banks")
    }
    for table, mapping in ids.items():
        print(f"  {table}: {len(mapping)}")

    def ref(table: str, name: str | None):
        if not name:
            return None
        key = name.strip()
        if key not in ids[table]:
            raise SystemExit(f"قيمة يتيمة في {table}: {name!r}")
        return ids[table][key]

    # 2) قيود الإيراد
    rev_rows = []
    for i, e in enumerate(revenue, start=1):
        movement = MOVEMENT_MAP.get(e["movement"] or "", "operational")
        rtype = e["revenue_type"]
        if movement == "operational" and not rtype:
            # الملف به ثلاثة قيود بلا نوع إيراد؛ نصنّفها بما يطابق بيانها
            rtype = (
                OPENING_REVENUE_TYPE
                if e["description"] and "افتتاحي" in e["description"]
                else UNKNOWN_REVENUE_TYPE
            )
        rev_rows.append(
            {
                "kind": "revenue",
                "entry_code": f"AUG-R-{i:04d}",
                "amount": e["amount"],
                "description": e["description"],
                "entry_date": str(e["entry_date"] or DEFAULT_REVENUE_DATE),
                "revenue_type_id": ref("revenue_types", rtype),
                "bank_id": ref("banks", e["bank"]),
                "movement_type": movement,
                "recon_status": RECON_MAP.get(e["recon"] or "", "pending"),
                "note": e["note"],
            }
        )

    # 3) قيود المصروف
    exp_rows = []
    for i, e in enumerate(expense, start=1):
        exp_rows.append(
            {
                "kind": "expense",
                "entry_code": e["entry_code"] or f"AUG-{i:04d}",
                "amount": e["amount"],
                "description": e["description"],
                "entry_date": str(e["entry_date"] or DEFAULT_REVENUE_DATE),
                "account_id": ref("accounts", e["account"]),
                "ledger_id": ref("ledgers", e["ledger"]),
                "cost_center_id": ref("cost_centers", e["cost_center"]),
                "bank_id": ref("banks", e["bank"]),
                "movement_type": MOVEMENT_MAP.get(e["movement"] or "", "operational"),
                "recon_status": RECON_MAP.get(e["recon"] or "", "pending"),
                "note": e["note"],
                "document_url": e["document"],
            }
        )

    existing = api.select("entries", "select=id&limit=1")
    if existing:
        print("  الجدول به قيود بالفعل — أحذفها قبل إعادة الترحيل")
        api.delete_all("entries")

    api.insert("entries", rev_rows)
    api.insert("entries", exp_rows)
    print(f"  رُحّل {len(rev_rows)} إيراد و{len(exp_rows)} مصروف")

    return sum(r["amount"] for r in rev_rows), sum(r["amount"] for r in exp_rows)


# --------------------------------------------------------------------------- verify

EXCEL_REVENUE = 75073854.12
EXCEL_EXPENSE = 78204322.30


def verify(api: Api):
    rows = api.rpc("f_operational_totals", {"p_from": None, "p_to": None})
    totals = rows[0] if isinstance(rows, list) else rows
    counts = api.select("entries", "select=kind")
    n_rev = sum(1 for r in counts if r["kind"] == "revenue")
    n_exp = sum(1 for r in counts if r["kind"] == "expense")

    # الإجمالي في الإكسل يشمل التحويلات الداخلية لأنه مجموع عمود المبلغ كاملًا
    all_rows = api.select("entries", "select=kind,amount")
    raw_rev = round(sum(float(r["amount"]) for r in all_rows if r["kind"] == "revenue"), 2)
    raw_exp = round(sum(float(r["amount"]) for r in all_rows if r["kind"] == "expense"), 2)

    print("\n== مطابقة الإجماليات مع الإكسل ==")
    print(f"عدد قيود الإيراد   : {n_rev:>6}  (الإكسل 30)")
    print(f"عدد قيود المصروف   : {n_exp:>6}  (الإكسل 206)")
    print(f"مجموع الإيراد      : {raw_rev:>18,.2f}  (الإكسل {EXCEL_REVENUE:,.2f})")
    print(f"مجموع المصروف      : {raw_exp:>18,.2f}  (الإكسل {EXCEL_EXPENSE:,.2f})")
    print(f"الإيراد التشغيلي   : {float(totals['revenue']):>18,.2f}")
    print(f"المصروف التشغيلي   : {float(totals['expense']):>18,.2f}")
    print(f"صافي التشغيل       : {float(totals['net']):>18,.2f}")

    cash = api.select("v_cash_position")[0]
    print("\n== المتاح فعليًا ==")
    print(f"أرصدة البنوك       : {float(cash['bank_total']):>18,.2f}")
    print(f"العهد القائمة      : {float(cash['custody_total']):>18,.2f}")
    print(f"إجمالي النقدية     : {float(cash['total_cash']):>18,.2f}")
    print(f"المديونيات المتبقية: {float(cash['debts_outstanding']):>18,.2f}")
    print(f"المتاح بعد السداد  : {float(cash['available_after_debts']):>18,.2f}")

    ok = (
        n_rev == 30
        and n_exp == 206
        and abs(raw_rev - EXCEL_REVENUE) < 0.05
        and abs(raw_exp - EXCEL_EXPENSE) < 0.05
    )
    print("\nالنتيجة:", "مطابق تمامًا ✔" if ok else "غير مطابق ✘")
    return ok


def main():
    require()
    api = Api(SUPABASE_URL, ANON_KEY)
    api.login(EMAIL, PASSWORD)
    print(f"تم الدخول كـ {EMAIL}")

    if "--verify" not in sys.argv:
        migrate(api)
    ok = verify(api)
    sys.exit(0 if ok else 1)


if __name__ == "__main__":
    main()
